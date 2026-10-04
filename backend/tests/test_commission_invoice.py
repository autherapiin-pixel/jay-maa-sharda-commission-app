"""Backend tests for Commission Invoice App.
Covers: settings, parties CRUD + search, invoices CRUD + math, exports (xlsx magic bytes),
stats, and the no-_id leak invariant.
"""
import os
import pytest
import requests

BASE_URL = os.environ.get("EXPO_PUBLIC_BACKEND_URL") or "https://drive-sync-billing.preview.emergentagent.com"
BASE_URL = BASE_URL.rstrip("/")
API = f"{BASE_URL}/api"

XLSX_MAGIC = b"PK\x03\x04"


# --------- fixtures ---------
@pytest.fixture(scope="session")
def session():
    s = requests.Session()
    s.headers.update({"Content-Type": "application/json"})
    return s


@pytest.fixture(scope="session")
def party(session):
    r = session.post(f"{API}/parties", json={
        "name": "TEST_Sardar Fruit Co",
        "location": "Prayagraj",
        "phone": "9999999999",
        "default_commission_percent": 6.0,
    })
    assert r.status_code == 200, r.text
    p = r.json()
    yield p
    # cleanup
    session.delete(f"{API}/parties/{p['id']}")


def _assert_no_underscore_id(obj):
    if isinstance(obj, dict):
        assert "_id" not in obj, f"_id leaked in response: {obj}"
        for v in obj.values():
            _assert_no_underscore_id(v)
    elif isinstance(obj, list):
        for v in obj:
            _assert_no_underscore_id(v)


# --------- settings ---------
class TestSettings:
    def test_get_defaults(self, session):
        r = session.get(f"{API}/settings")
        assert r.status_code == 200
        data = r.json()
        _assert_no_underscore_id(data)
        for key in ("shop_name", "shop_tagline", "shop_address", "default_commission_percent"):
            assert key in data
        assert isinstance(data["default_commission_percent"], (int, float))

    def test_update_settings(self, session):
        payload = {
            "shop_name": "TEST_Shop",
            "shop_tagline": "TEST_Tagline",
            "shop_address": "TEST_Address",
            "default_commission_percent": 7.5,
        }
        r = session.put(f"{API}/settings", json=payload)
        assert r.status_code == 200
        data = r.json()
        assert data["shop_name"] == "TEST_Shop"
        assert data["default_commission_percent"] == 7.5
        # verify persisted
        r2 = session.get(f"{API}/settings")
        assert r2.json()["shop_name"] == "TEST_Shop"
        # restore
        session.put(f"{API}/settings", json={
            "shop_name": "जय माँ शारदा फ्रूट सेंटर",
            "shop_tagline": "फ्रूट के थोक विक्रेता एवं कमीशन एजेंट",
            "shop_address": "नवीन फल मंडी, सिरसा, मेजारोड - प्रयागराज",
            "default_commission_percent": 6.0,
        })


# --------- parties ---------
class TestParties:
    def test_create_list_search(self, session):
        r = session.post(f"{API}/parties", json={"name": "TEST_SearchableAlpha"})
        assert r.status_code == 200
        pid = r.json()["id"]
        try:
            # list
            lr = session.get(f"{API}/parties")
            assert lr.status_code == 200
            names = [p["name"] for p in lr.json()]
            assert "TEST_SearchableAlpha" in names
            _assert_no_underscore_id(lr.json())
            # search case-insensitive
            sr = session.get(f"{API}/parties", params={"q": "searchablealpha"})
            assert sr.status_code == 200
            assert any(p["id"] == pid for p in sr.json())
        finally:
            session.delete(f"{API}/parties/{pid}")

    def test_update_delete(self, session):
        r = session.post(f"{API}/parties", json={"name": "TEST_PartyUpd"})
        pid = r.json()["id"]
        upd = session.put(f"{API}/parties/{pid}", json={"name": "TEST_PartyUpd2", "phone": "123"})
        assert upd.status_code == 200
        assert upd.json()["name"] == "TEST_PartyUpd2"
        get_r = session.get(f"{API}/parties/{pid}")
        assert get_r.json()["phone"] == "123"
        d = session.delete(f"{API}/parties/{pid}")
        assert d.status_code == 200
        # confirm gone
        assert session.get(f"{API}/parties/{pid}").status_code == 404


# --------- invoices math + CRUD ---------
class TestInvoices:
    def test_reference_math(self, session, party):
        items = [
            {"sr": 10, "item": "Apple", "qty": 100, "rate": 220, "total": 22000},
            {"sr": 5, "item": "Banana", "qty": 50, "rate": 50, "total": 2500},
            {"sr": 15, "item": "Mango", "qty": 70, "rate": 315, "total": 22050},
        ]
        payload = {
            "party_id": party["id"],
            "date": "2025-01-15",
            "items": items,
            "bhada": 1700,
            "mazdoori": 225,
            "commission_percent": 6,
            "sr_total": 30,
        }
        r = session.post(f"{API}/invoices", json=payload)
        assert r.status_code == 200, r.text
        inv = r.json()
        _assert_no_underscore_id(inv)
        assert inv["items_total"] == 46550
        assert inv["commission"] == 2793
        assert inv["kharch_total"] == 4718
        assert inv["net"] == 41832
        # verify get
        g = session.get(f"{API}/invoices/{inv['id']}")
        assert g.status_code == 200
        assert g.json()["net"] == 41832
        # verify filter by party
        lp = session.get(f"{API}/invoices", params={"party_id": party["id"]})
        assert any(x["id"] == inv["id"] for x in lp.json())
        # update: change bhada and recompute
        up = session.put(f"{API}/invoices/{inv['id']}", json={"bhada": 2000})
        assert up.status_code == 200
        upd = up.json()
        # kharch = 2000 + 225 + 2793 = 5018; net = 46550 - 5018 = 41532
        assert upd["kharch_total"] == 5018
        assert upd["net"] == 41532
        # export xlsx
        ex = session.get(f"{API}/invoices/{inv['id']}/export")
        assert ex.status_code == 200
        assert ex.content[:4] == XLSX_MAGIC
        assert "spreadsheetml" in ex.headers.get("Content-Type", "")
        # delete
        d = session.delete(f"{API}/invoices/{inv['id']}")
        assert d.status_code == 200
        assert session.get(f"{API}/invoices/{inv['id']}").status_code == 404

    def test_invalid_party(self, session):
        r = session.post(f"{API}/invoices", json={
            "party_id": "does-not-exist",
            "date": "2025-01-15",
            "items": [],
            "bhada": 0, "mazdoori": 0, "commission_percent": 6,
        })
        assert r.status_code == 400


# --------- party export + stats ---------
class TestPartyExportAndStats:
    def test_party_export_multisheet(self, session, party):
        # create two invoices
        ids = []
        for d, total in [("2025-01-10", 1000), ("2025-01-12", 2000)]:
            r = session.post(f"{API}/invoices", json={
                "party_id": party["id"],
                "date": d,
                "items": [{"sr": 1, "item": "X", "qty": 1, "rate": total, "total": total}],
                "bhada": 0, "mazdoori": 0, "commission_percent": 6, "sr_total": 1,
            })
            assert r.status_code == 200
            ids.append(r.json()["id"])
        try:
            ex = session.get(f"{API}/parties/{party['id']}/export")
            assert ex.status_code == 200
            assert ex.content[:4] == XLSX_MAGIC
            # verify multisheet using openpyxl
            from openpyxl import load_workbook
            from io import BytesIO
            wb = load_workbook(BytesIO(ex.content))
            assert len(wb.sheetnames) >= 2
        finally:
            for i in ids:
                session.delete(f"{API}/invoices/{i}")

    def test_stats_shape(self, session):
        r = session.get(f"{API}/stats")
        assert r.status_code == 200
        data = r.json()
        for key in ("parties", "invoices", "total_net", "total_commission", "recent"):
            assert key in data
        assert isinstance(data["recent"], list)
        _assert_no_underscore_id(data)
