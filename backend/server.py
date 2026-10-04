from fastapi import FastAPI, APIRouter, HTTPException, Depends, File, UploadFile
from fastapi.concurrency import run_in_threadpool
from fastapi.responses import StreamingResponse, Response
from starlette.middleware.cors import CORSMiddleware
import os
import logging
import re
import hashlib
from pydantic import BaseModel, Field
from typing import List, Optional
import uuid
from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from zipfile import BadZipFile

from database import db, client  # loads .env first
import google_sync
import auth as auth_mod
import storage as storage_mod
from auth import get_current_user
from pdf_gen import build_invoice_pdf
from excel_import import parse_commission_workbook


def _safe_filename(s: str) -> str:
    # Replace all non-ASCII with underscore so Content-Disposition stays latin-1 safe.
    ascii_ = re.sub(r"[^A-Za-z0-9._-]+", "_", s).strip("_")
    return ascii_ or "invoice"


def _uid(user: dict) -> str:
    return user["user_id"]


app = FastAPI()
api_router = APIRouter(prefix="/api")
public_router = APIRouter(prefix="/api")


# ---------------- MODELS ----------------
class Settings(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: Optional[str] = None
    shop_name: str = "जय माँ शारदा फ्रूट सेंटर"
    shop_tagline: str = "फ्रूट के थोक विक्रेता एवं कमीशन एजेंट"
    shop_address: str = "नवीन फल मंडी, सिरसा, मेजारोड - प्रयागराज"
    default_commission_percent: float = 6.0
    whatsapp_number: Optional[str] = ""
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class SettingsUpdate(BaseModel):
    shop_name: Optional[str] = None
    shop_tagline: Optional[str] = None
    shop_address: Optional[str] = None
    default_commission_percent: Optional[float] = None
    whatsapp_number: Optional[str] = None


class Party(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: Optional[str] = None
    name: str
    location: Optional[str] = ""
    phone: Optional[str] = ""
    default_commission_percent: Optional[float] = None
    drive_folder_id: Optional[str] = None
    drive_folder_url: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class PartyCreate(BaseModel):
    name: str
    location: Optional[str] = ""
    phone: Optional[str] = ""
    default_commission_percent: Optional[float] = None


class PartyUpdate(BaseModel):
    name: Optional[str] = None
    location: Optional[str] = None
    phone: Optional[str] = None
    default_commission_percent: Optional[float] = None


class InvoiceItem(BaseModel):
    sr: Optional[int] = None  # free SR number (user types, like crate count)
    item: str
    qty: float = 0
    rate: float = 0
    total: float = 0  # qty * rate (also user-editable to match legacy pattern)


class Payment(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    amount: float
    date: str  # yyyy-mm-dd
    mode: str = "cash"  # cash | upi | bank | cheque
    note: Optional[str] = ""
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class PaymentCreate(BaseModel):
    amount: float
    date: Optional[str] = None
    mode: str = "cash"
    note: Optional[str] = ""


class Invoice(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    user_id: Optional[str] = None
    party_id: str
    party_name: str
    date: str  # ISO date yyyy-mm-dd
    items: List[InvoiceItem] = []
    sr_total: Optional[float] = 0  # bottom-left crate count sum
    items_total: float = 0  # sum of row totals
    bhada: float = 0
    mazdoori: float = 0
    commission_percent: float = 6.0
    commission: float = 0
    kharch_total: float = 0  # bhada + mazdoori + commission
    net: float = 0  # items_total - kharch_total
    note: Optional[str] = ""
    description: Optional[str] = ""
    photo_path: Optional[str] = None
    payments: List[Payment] = []
    paid_amount: float = 0  # sum of payments (payment out to party)
    balance: float = 0  # net - paid_amount
    paid: bool = False  # derived: balance <= 0
    paid_at: Optional[datetime] = None
    drive_file_id: Optional[str] = None
    drive_file_url: Optional[str] = None
    drive_pdf_id: Optional[str] = None
    drive_pdf_url: Optional[str] = None
    synced_at: Optional[datetime] = None
    import_source_key: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class InvoiceCreate(BaseModel):
    party_id: str
    date: str
    items: List[InvoiceItem] = []
    bhada: float = 0
    mazdoori: float = 0
    commission_percent: float = 6.0
    sr_total: Optional[float] = 0
    note: Optional[str] = ""
    description: Optional[str] = ""
    photo_path: Optional[str] = None
    payment_out: float = 0  # paid to party at bill time
    payment_out_mode: str = "cash"


class InvoiceUpdate(BaseModel):
    date: Optional[str] = None
    items: Optional[List[InvoiceItem]] = None
    bhada: Optional[float] = None
    mazdoori: Optional[float] = None
    commission_percent: Optional[float] = None
    sr_total: Optional[float] = None
    note: Optional[str] = None
    description: Optional[str] = None
    photo_path: Optional[str] = None


# ---------------- HELPERS ----------------
def compute_totals(items: List[InvoiceItem], bhada: float, mazdoori: float, commission_percent: float):
    items_total = 0.0
    for it in items:
        # if total not set but qty*rate present, derive it
        if (not it.total) and it.qty and it.rate:
            it.total = round(it.qty * it.rate, 2)
        items_total += it.total or 0
    commission = round(items_total * (commission_percent or 0) / 100, 2)
    kharch_total = round((bhada or 0) + (mazdoori or 0) + commission, 2)
    net = round(items_total - kharch_total, 2)
    return items_total, commission, kharch_total, net


def apply_payments(data: dict) -> dict:
    """Derive paid_amount / balance / paid from the payments list (mutates and returns data)."""
    payments = data.get("payments") or []
    paid_amount = round(sum((p.get("amount") or 0) for p in payments), 2)
    net = data.get("net") or 0
    balance = round(net - paid_amount, 2)
    data["payments"] = payments
    data["paid_amount"] = paid_amount
    data["balance"] = balance
    data["paid"] = balance <= 0.005
    if data["paid"]:
        if not data.get("paid_at"):
            data["paid_at"] = datetime.now(timezone.utc)
    else:
        data["paid_at"] = None
    return data


async def migrate_payments():
    """One-time: convert legacy paid flag into a payment entry for docs without `payments`."""
    cursor = db.invoices.find({"payments": {"$exists": False}}, {"_id": 0})
    async for inv in cursor:
        payments = []
        if inv.get("paid") and (inv.get("net") or 0) > 0:
            paid_at = inv.get("paid_at")
            pay_date = paid_at.strftime("%Y-%m-%d") if isinstance(paid_at, datetime) else inv.get("date", "")
            payments.append(Payment(amount=inv["net"], date=pay_date, mode="cash", note="Paid (migrated)").model_dump())
        inv["payments"] = payments
        apply_payments(inv)
        await db.invoices.update_one(
            {"id": inv["id"]},
            {"$set": {k: inv[k] for k in ("payments", "paid_amount", "balance", "paid", "paid_at")}},
        )


# Mongo expressions reused by aggregations
_PAID_AMT = {"$ifNull": ["$paid_amount", 0]}
_BALANCE = {"$max": [{"$subtract": ["$net", _PAID_AMT]}, 0]}
_IS_PENDING = {"$gt": [{"$subtract": ["$net", _PAID_AMT]}, 0.005]}


def clean(doc: dict) -> dict:
    if not doc:
        return doc
    doc.pop("_id", None)
    return doc


# ---------------- SETTINGS ----------------
async def _get_settings_doc(uid: str) -> dict:
    doc = await db.settings.find_one({"user_id": uid}, {"_id": 0})
    if not doc:
        s = Settings(user_id=uid)
        await db.settings.insert_one(s.model_dump())
        return s.model_dump()
    return doc


@api_router.get("/settings", response_model=Settings)
async def get_settings(user: dict = Depends(get_current_user)):
    return Settings(**await _get_settings_doc(_uid(user)))


@api_router.put("/settings", response_model=Settings)
async def update_settings(upd: SettingsUpdate, user: dict = Depends(get_current_user)):
    data = await _get_settings_doc(_uid(user))
    for k, v in upd.model_dump(exclude_unset=True).items():
        if v is not None:
            data[k] = v
    data["updated_at"] = datetime.now(timezone.utc)
    await db.settings.update_one({"user_id": _uid(user)}, {"$set": data}, upsert=True)
    return Settings(**data)


# ---------------- PARTIES ----------------
@api_router.get("/parties", response_model=List[Party])
async def list_parties(q: Optional[str] = None, user: dict = Depends(get_current_user)):
    query: dict = {"user_id": _uid(user)}
    if q:
        query["name"] = {"$regex": q, "$options": "i"}
    docs = await db.parties.find(query, {"_id": 0}).sort("name", 1).to_list(1000)
    return [Party(**d) for d in docs]


@api_router.post("/parties", response_model=Party)
async def create_party(p: PartyCreate, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    party = Party(**p.model_dump(), user_id=uid)
    await db.parties.insert_one(party.model_dump())
    # Create Google Drive folder if connected
    try:
        if await google_sync.is_connected(db, uid):
            await google_sync.ensure_party_folder(db, uid, party.model_dump())
            fresh = await db.parties.find_one({"id": party.id}, {"_id": 0})
            if fresh:
                return Party(**fresh)
    except Exception as e:
        logging.getLogger(__name__).warning("drive folder create skipped: %s", e)
    return party


@api_router.get("/parties/{party_id}", response_model=Party)
async def get_party(party_id: str, user: dict = Depends(get_current_user)):
    doc = await db.parties.find_one({"id": party_id, "user_id": _uid(user)}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Party not found")
    return Party(**doc)


@api_router.put("/parties/{party_id}", response_model=Party)
async def update_party(party_id: str, upd: PartyUpdate, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    data = {k: v for k, v in upd.model_dump(exclude_unset=True).items() if v is not None}
    if data:
        await db.parties.update_one({"id": party_id, "user_id": uid}, {"$set": data})
    doc = await db.parties.find_one({"id": party_id, "user_id": uid}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Party not found")
    # keep invoice party_name snapshots in sync if name changed
    if "name" in data:
        await db.invoices.update_many({"party_id": party_id}, {"$set": {"party_name": data["name"]}})
    return Party(**doc)


@api_router.delete("/parties/{party_id}")
async def delete_party(party_id: str, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    # Collect drive refs before delete so we can clean the sheet rows
    invoices = await db.invoices.find({"party_id": party_id, "user_id": uid}, {"_id": 0}).to_list(10000)
    await db.parties.delete_one({"id": party_id, "user_id": uid})
    await db.invoices.delete_many({"party_id": party_id, "user_id": uid})
    # Remove sheet rows (best-effort)
    try:
        for inv in invoices:
            await google_sync.remove_master_sheet_row(db, uid, inv["id"])
        await google_sync.sync_ledger(db, uid)
    except Exception:
        pass
    return {"ok": True}


# ---------------- INVOICES ----------------
@api_router.get("/invoices", response_model=List[Invoice])
async def list_invoices(party_id: Optional[str] = None, limit: int = 500, user: dict = Depends(get_current_user)):
    query: dict = {"user_id": _uid(user)}
    if party_id:
        query["party_id"] = party_id
    docs = await db.invoices.find(query, {"_id": 0}).sort("date", -1).to_list(limit)
    return [Invoice(**d) for d in docs]


@api_router.post("/invoices", response_model=Invoice)
async def create_invoice(inv: InvoiceCreate, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    party = await db.parties.find_one({"id": inv.party_id, "user_id": uid}, {"_id": 0})
    if not party:
        raise HTTPException(400, "Party not found")
    items_total, commission, kharch_total, net = compute_totals(
        inv.items, inv.bhada, inv.mazdoori, inv.commission_percent
    )
    payments = []
    if inv.payment_out and inv.payment_out > 0:
        payments.append(Payment(amount=round(inv.payment_out, 2), date=inv.date,
                                mode=inv.payment_out_mode or "cash", note="बिल के समय दिया"))
    invoice = Invoice(
        user_id=uid,
        party_id=inv.party_id,
        party_name=party["name"],
        date=inv.date,
        items=inv.items,
        sr_total=inv.sr_total or 0,
        items_total=items_total,
        bhada=inv.bhada,
        mazdoori=inv.mazdoori,
        commission_percent=inv.commission_percent,
        commission=commission,
        kharch_total=kharch_total,
        net=net,
        note=inv.note or "",
        description=inv.description or "",
        photo_path=inv.photo_path,
        payments=payments,
    )
    doc = apply_payments(invoice.model_dump())
    await db.invoices.insert_one(doc)
    await _sync_invoice_to_google(uid, invoice.id)
    fresh = await db.invoices.find_one({"id": invoice.id}, {"_id": 0})
    return Invoice(**(fresh or invoice.model_dump()))


@api_router.get("/invoices/{invoice_id}", response_model=Invoice)
async def get_invoice(invoice_id: str, user: dict = Depends(get_current_user)):
    doc = await db.invoices.find_one({"id": invoice_id, "user_id": _uid(user)}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Invoice not found")
    return Invoice(**doc)


@api_router.put("/invoices/{invoice_id}", response_model=Invoice)
async def update_invoice(invoice_id: str, upd: InvoiceUpdate, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    existing = await db.invoices.find_one({"id": invoice_id, "user_id": uid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Invoice not found")
    data = existing.copy()
    for k, v in upd.model_dump(exclude_unset=True).items():
        if v is not None or k == "photo_path":
            data[k] = v
    items = [InvoiceItem(**(i if isinstance(i, dict) else i.model_dump())) for i in data.get("items", [])]
    items_total, commission, kharch_total, net = compute_totals(
        items, data.get("bhada", 0), data.get("mazdoori", 0), data.get("commission_percent", 6.0)
    )
    data["items"] = [i.model_dump() for i in items]
    data["items_total"] = items_total
    data["commission"] = commission
    data["kharch_total"] = kharch_total
    data["net"] = net
    apply_payments(data)
    data["updated_at"] = datetime.now(timezone.utc)
    await db.invoices.update_one({"id": invoice_id}, {"$set": data})
    doc = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
    await _sync_invoice_to_google(uid, invoice_id)
    return Invoice(**doc)


@api_router.delete("/invoices/{invoice_id}")
async def delete_invoice(invoice_id: str, user: dict = Depends(get_current_user)):
    uid = _uid(user)
    existing = await db.invoices.find_one({"id": invoice_id, "user_id": uid}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Invoice not found")
    await db.invoices.delete_one({"id": invoice_id})
    try:
        for fid in (existing.get("drive_file_id"), existing.get("drive_pdf_id")):
            if fid:
                await google_sync.delete_drive_file(db, uid, fid)
        await google_sync.remove_master_sheet_row(db, uid, invoice_id)
        await google_sync.sync_ledger(db, uid)
    except Exception:
        pass
    return {"ok": True}


class PaymentUpdate(BaseModel):
    paid: bool


async def _find_invoice(invoice_id: str, user: dict) -> dict:
    existing = await db.invoices.find_one({"id": invoice_id, "user_id": _uid(user)}, {"_id": 0})
    if not existing:
        raise HTTPException(404, "Invoice not found")
    return existing


async def _save_payments(uid: str, invoice_id: str, existing: dict, payments: List[dict]) -> Invoice:
    existing["payments"] = payments
    apply_payments(existing)
    patch = {k: existing[k] for k in ("payments", "paid_amount", "balance", "paid", "paid_at")}
    patch["updated_at"] = datetime.now(timezone.utc)
    await db.invoices.update_one({"id": invoice_id}, {"$set": patch})
    doc = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
    try:
        await _sync_invoice_to_google(uid, invoice_id)
    except Exception:
        pass
    return Invoice(**doc)


@api_router.put("/invoices/{invoice_id}/payment", response_model=Invoice)
async def toggle_payment(invoice_id: str, upd: PaymentUpdate, user: dict = Depends(get_current_user)):
    """Quick action: paid=True records a payment for the remaining balance; paid=False clears all payments."""
    existing = await _find_invoice(invoice_id, user)
    payments = list(existing.get("payments") or [])
    if upd.paid:
        balance = round((existing.get("net") or 0) - sum(p.get("amount") or 0 for p in payments), 2)
        if balance > 0:
            payments.append(Payment(
                amount=balance,
                date=datetime.now(timezone.utc).strftime("%Y-%m-%d"),
                mode="cash",
                note="पूरा Paid",
            ).model_dump())
    else:
        payments = []
    return await _save_payments(_uid(user), invoice_id, existing, payments)


@api_router.post("/invoices/{invoice_id}/payments", response_model=Invoice)
async def add_payment(invoice_id: str, body: PaymentCreate, user: dict = Depends(get_current_user)):
    if body.amount <= 0:
        raise HTTPException(400, "Amount must be > 0")
    existing = await _find_invoice(invoice_id, user)
    payments = list(existing.get("payments") or [])
    payments.append(Payment(
        amount=round(body.amount, 2),
        date=body.date or datetime.now(timezone.utc).strftime("%Y-%m-%d"),
        mode=body.mode or "cash",
        note=body.note or "",
    ).model_dump())
    return await _save_payments(_uid(user), invoice_id, existing, payments)


@api_router.delete("/invoices/{invoice_id}/payments/{payment_id}", response_model=Invoice)
async def delete_payment(invoice_id: str, payment_id: str, user: dict = Depends(get_current_user)):
    existing = await _find_invoice(invoice_id, user)
    payments = [p for p in (existing.get("payments") or []) if p.get("id") != payment_id]
    if len(payments) == len(existing.get("payments") or []):
        raise HTTPException(404, "Payment not found")
    return await _save_payments(_uid(user), invoice_id, existing, payments)


@api_router.get("/parties/{party_id}/payments")
async def party_payments(party_id: str, user: dict = Depends(get_current_user)):
    """Flat list of all payments given to a party (newest first)."""
    docs = await db.invoices.find(
        {"party_id": party_id, "user_id": _uid(user), "payments.0": {"$exists": True}},
        {"_id": 0, "id": 1, "date": 1, "net": 1, "payments": 1},
    ).to_list(1000)
    out = []
    for d in docs:
        for p in d.get("payments") or []:
            out.append({**p, "invoice_id": d["id"], "invoice_date": d.get("date"), "invoice_net": d.get("net")})
    out.sort(key=lambda p: (p.get("date") or "", str(p.get("created_at") or "")), reverse=True)
    return out


@api_router.get("/parties/{party_id}/items")
async def party_items(party_id: str, user: dict = Depends(get_current_user)):
    """Return distinct items this party has used, newest first with last rate."""
    cursor = db.invoices.find(
        {"party_id": party_id, "user_id": _uid(user)}, {"_id": 0, "items": 1, "date": 1, "created_at": 1}
    ).sort("created_at", -1)
    docs = await cursor.to_list(1000)
    seen = {}
    for d in docs:
        for it in d.get("items") or []:
            name = (it.get("item") or "").strip()
            if not name:
                continue
            key = name.lower()
            if key in seen:
                seen[key]["count"] += 1
                continue
            seen[key] = {
                "item": name,
                "last_rate": it.get("rate") or 0,
                "last_qty": it.get("qty") or 0,
                "last_total": it.get("total") or 0,
                "count": 1,
            }
    return list(seen.values())


@api_router.get("/parties/{party_id}/summary")
async def party_summary(party_id: str, user: dict = Depends(get_current_user)):
    """Totals + pending (unpaid) counts for a party."""
    agg = await db.invoices.aggregate([
        {"$match": {"party_id": party_id, "user_id": _uid(user)}},
        {"$group": {
            "_id": None,
            "count": {"$sum": 1},
            "total_net": {"$sum": "$net"},
            "total_commission": {"$sum": "$commission"},
            "pending_count": {"$sum": {"$cond": [_IS_PENDING, 1, 0]}},
            "pending_net": {"$sum": _BALANCE},
            "paid_net": {"$sum": _PAID_AMT},
        }},
    ]).to_list(1)
    if not agg:
        return {"count": 0, "total_net": 0, "total_commission": 0,
                "pending_count": 0, "pending_net": 0, "paid_net": 0}
    r = agg[0]
    r.pop("_id", None)
    return r


@api_router.get("/stats/daybook")
async def daybook(period: str = "day", frm: Optional[str] = None, to: Optional[str] = None,
                  user: dict = Depends(get_current_user)):
    """
    period = 'day' or 'month'. frm/to are yyyy-mm-dd strings for filtering `date` field.
    """
    match: dict = {"user_id": _uid(user)}
    if frm:
        match["date"] = {"$gte": frm}
    if to:
        match.setdefault("date", {})["$lte"] = to

    # Group key: 'day' -> entire date; 'month' -> yyyy-mm
    group_id = (
        "$date"
        if period == "day"
        else {"$substr": ["$date", 0, 7]}  # yyyy-mm
    )

    pipeline = [{"$match": match}]
    pipeline += [
        {"$group": {
            "_id": group_id,
            "invoices": {"$sum": 1},
            "items_total": {"$sum": "$items_total"},
            "bhada": {"$sum": "$bhada"},
            "mazdoori": {"$sum": "$mazdoori"},
            "commission": {"$sum": "$commission"},
            "net": {"$sum": "$net"},
            "paid_net": {"$sum": _PAID_AMT},
            "pending_net": {"$sum": _BALANCE},
        }},
        {"$sort": {"_id": -1}},
    ]
    rows = await db.invoices.aggregate(pipeline).to_list(1000)
    out = []
    for r in rows:
        r["period"] = r.pop("_id")
        out.append(r)
    # overall totals
    totals = {k: sum(r.get(k, 0) for r in out) for k in
              ["invoices", "items_total", "bhada", "mazdoori",
               "commission", "net", "paid_net", "pending_net"]}
    return {"rows": out, "totals": totals, "period": period}


# ---------------- STATS ----------------
@api_router.get("/stats")
async def stats(user: dict = Depends(get_current_user)):
    uid = _uid(user)
    parties_count = await db.parties.count_documents({"user_id": uid})
    invoices_count = await db.invoices.count_documents({"user_id": uid})
    pipeline = [{"$match": {"user_id": uid}}, {"$group": {
        "_id": None,
        "total_net": {"$sum": "$net"},
        "total_commission": {"$sum": "$commission"},
        "pending_net": {"$sum": _BALANCE},
        "pending_count": {"$sum": {"$cond": [_IS_PENDING, 1, 0]}},
    }}]
    agg = await db.invoices.aggregate(pipeline).to_list(1)
    base = {"total_net": 0, "total_commission": 0, "pending_net": 0, "pending_count": 0}
    if agg:
        a = agg[0]
        base = {k: a.get(k, 0) for k in base}
    recent = await db.invoices.find({"user_id": uid}, {"_id": 0}).sort("created_at", -1).limit(5).to_list(5)
    return {
        "parties": parties_count,
        "invoices": invoices_count,
        **base,
        "recent": recent,
    }


# ---------------- EXCEL IMPORT / EXPORT ----------------
@api_router.post("/import/excel")
async def import_excel(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Please choose an .xlsx workbook")

    contents = await file.read(20 * 1024 * 1024 + 1)
    if len(contents) > 20 * 1024 * 1024:
        raise HTTPException(413, "Workbook is larger than the 20 MB import limit")
    if not contents:
        raise HTTPException(400, "Workbook is empty")

    try:
        invoices, warnings = await run_in_threadpool(parse_commission_workbook, contents)
    except (ValueError, OSError) as exc:
        raise HTTPException(400, f"Could not read workbook: {exc}")
    except (InvalidFileException, BadZipFile, EOFError):
        raise HTTPException(400, "The selected file is not a valid .xlsx workbook")

    uid = _uid(user)
    parties = await db.parties.find({"user_id": uid}, {"_id": 0}).to_list(10000)
    party_by_name = {p["name"].strip().casefold(): p for p in parties}
    imported = 0
    skipped = 0
    parties_created = 0

    for record in invoices:
        party_name = record["party_name"].strip()
        party_key = party_name.casefold()
        party = party_by_name.get(party_key)
        if not party:
            party = Party(name=party_name, user_id=uid)
            party_doc = party.model_dump()
            await db.parties.insert_one(party_doc)
            party_by_name[party_key] = party_doc
            parties_created += 1

        source_identity = "\x1f".join((
            uid, record["sheet_name"], str(record["source_row"]),
            record["date"], party_name.casefold(),
        ))
        source_key = hashlib.sha256(source_identity.encode("utf-8")).hexdigest()
        items = [InvoiceItem(**item) for item in record["items"]]
        items_total, commission, kharch_total, net = compute_totals(
            items, record["bhada"], record["mazdoori"], record["commission_percent"]
        )
        invoice = Invoice(
            user_id=uid,
            party_id=party["id"],
            party_name=party_name,
            date=record["date"],
            items=items,
            sr_total=record["sr_total"],
            items_total=items_total,
            bhada=record["bhada"],
            mazdoori=record["mazdoori"],
            commission_percent=record["commission_percent"],
            commission=commission,
            kharch_total=kharch_total,
            net=net,
            import_source_key=source_key,
        )
        doc = apply_payments(invoice.model_dump())
        result = await db.invoices.update_one(
            {"user_id": uid, "import_source_key": source_key},
            {"$setOnInsert": doc},
            upsert=True,
        )
        if result.upserted_id is None:
            skipped += 1
        else:
            imported += 1

    return {
        "parties_created": parties_created,
        "invoices_imported": imported,
        "invoices_skipped": skipped,
        "warnings": warnings,
    }


def _build_invoice_sheet(ws, invoice: dict, settings: dict):
    thin = Side(border_style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Column widths roughly matching the paper invoice
    widths = [8, 24, 10, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Header orange band
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = settings.get("shop_name", "")
    c.fill = PatternFill("solid", fgColor="F59E0B")
    c.font = Font(bold=True, size=16, color="000000")
    c.alignment = center
    ws.row_dimensions[1].height = 32

    # Green tagline band
    ws.merge_cells("A2:E2")
    c = ws["A2"]
    c.value = settings.get("shop_tagline", "")
    c.fill = PatternFill("solid", fgColor="86C33B")
    c.font = Font(bold=True, size=12)
    c.alignment = center
    ws.row_dimensions[2].height = 24

    # Dark address band
    ws.merge_cells("A3:E3")
    c = ws["A3"]
    c.value = settings.get("shop_address", "")
    c.fill = PatternFill("solid", fgColor="1F1F1F")
    c.font = Font(bold=True, color="FFFFFF")
    c.alignment = center
    ws.row_dimensions[3].height = 22

    # Party + date
    ws["A4"] = "पार्टी नाम:"
    ws["A4"].font = Font(bold=True)
    ws.merge_cells("B4:C4")
    ws["B4"] = invoice.get("party_name", "")
    ws["B4"].font = Font(bold=True)
    ws["D4"] = "Date:"
    ws["D4"].font = Font(bold=True)
    ws["E4"] = invoice.get("date", "")
    for cell in ["A4", "B4", "D4", "E4"]:
        ws[cell].border = border
        ws[cell].alignment = left

    # Table head
    heads = ["SR", "ITEM", "QTY", "RATE", "TOTAL"]
    for i, h in enumerate(heads, start=1):
        cc = ws.cell(row=5, column=i, value=h)
        cc.fill = PatternFill("solid", fgColor="111111")
        cc.font = Font(bold=True, color="FFFFFF")
        cc.alignment = center
        cc.border = border
    ws.row_dimensions[5].height = 22

    # Items (pad to 10 rows minimum for the printed look)
    items = invoice.get("items", [])
    row_count = max(10, len(items))
    r = 6
    for i in range(row_count):
        it = items[i] if i < len(items) else {}
        values = [
            it.get("sr") if it else "",
            it.get("item", "") if it else "",
            it.get("qty") if it else "",
            it.get("rate") if it else "",
            it.get("total") if it else "",
        ]
        for col, v in enumerate(values, start=1):
            cc = ws.cell(row=r, column=col, value=v if v not in (None, 0) else (v if v == 0 and i < len(items) else ""))
            cc.border = border
            cc.alignment = center if col != 2 else left
            if col == 5:
                cc.fill = PatternFill("solid", fgColor="CFE7F6")
        r += 1

    # SR total + TOTAL row
    ws.cell(row=r, column=1, value=invoice.get("sr_total") or "").fill = PatternFill("solid", fgColor="F4B678")
    ws.cell(row=r, column=1).border = border
    ws.cell(row=r, column=1).alignment = center
    ws.cell(row=r, column=1).font = Font(bold=True)
    for col in (2, 3):
        cc = ws.cell(row=r, column=col, value="")
        cc.border = border
        cc.fill = PatternFill("solid", fgColor="F4B678")
    tc = ws.cell(row=r, column=4, value="TOTAL")
    tc.fill = PatternFill("solid", fgColor="F4B678")
    tc.font = Font(bold=True)
    tc.alignment = center
    tc.border = border
    vc = ws.cell(row=r, column=5, value=invoice.get("items_total", 0))
    vc.fill = PatternFill("solid", fgColor="CFE7F6")
    vc.font = Font(bold=True)
    vc.alignment = center
    vc.border = border
    r += 2

    # Expense block
    pct = invoice.get("commission_percent", 6.0)
    rows = [
        ("भाड़ा", invoice.get("bhada", 0)),
        ("मजदूरी", invoice.get("mazdoori", 0)),
        (f"कमीशन ({pct:g}%)", invoice.get("commission", 0)),
        ("खर्च टोटल", invoice.get("kharch_total", 0)),
    ]
    for label, val in rows:
        lc = ws.cell(row=r, column=1, value=label)
        lc.fill = PatternFill("solid", fgColor="DCE9F7")
        lc.font = Font(bold=True)
        lc.border = border
        lc.alignment = left
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        vc = ws.cell(row=r, column=3, value=val)
        vc.border = border
        vc.alignment = center
        for col in (4, 5):
            ws.cell(row=r, column=col, value="").border = border
        r += 1

    # NET row — yellow
    ws.cell(row=r, column=1, value="").border = border
    ws.cell(row=r, column=2, value="").border = border
    ws.cell(row=r, column=3, value="").border = border
    nc = ws.cell(row=r, column=4, value="नेट")
    nc.fill = PatternFill("solid", fgColor="FFEF00")
    nc.font = Font(bold=True)
    nc.alignment = center
    nc.border = border
    vn = ws.cell(row=r, column=5, value=invoice.get("net", 0))
    vn.fill = PatternFill("solid", fgColor="FFEF00")
    vn.font = Font(bold=True)
    vn.alignment = center
    vn.border = border


@api_router.get("/invoices/{invoice_id}/export")
async def export_invoice_excel(invoice_id: str, user: dict = Depends(auth_mod.get_current_user_any)):
    inv = await _find_invoice(invoice_id, user)
    settings_doc = await _get_settings_doc(_uid(user))

    wb = Workbook()
    ws = wb.active
    ws.title = inv.get("date", "Invoice")[:31]
    _build_invoice_sheet(ws, inv, settings_doc)
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_party = _safe_filename(inv.get("party_name") or "party")
    filename = f"{safe_party}_{inv.get('date','')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


async def _invoice_pdf_bytes(inv: dict, settings_doc: dict) -> bytes:
    photo = None
    if inv.get("photo_path"):
        try:
            photo, _ = await storage_mod.run_in_threadpool(storage_mod._get, inv["photo_path"])
        except Exception as e:
            logger.warning("photo fetch for pdf failed: %s", e)
    return await storage_mod.run_in_threadpool(build_invoice_pdf, inv, settings_doc, photo)


@api_router.get("/invoices/{invoice_id}/pdf")
async def invoice_pdf(invoice_id: str, user: dict = Depends(auth_mod.get_current_user_any)):
    inv = await _find_invoice(invoice_id, user)
    settings_doc = await _get_settings_doc(_uid(user))
    pdf = await _invoice_pdf_bytes(inv, settings_doc)
    filename = f"{_safe_filename(inv.get('party_name') or 'party')}_{inv.get('date','')}.pdf"
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename={filename}"})


@public_router.get("/public/invoices/{invoice_id}/pdf")
async def public_invoice_pdf(invoice_id: str):
    """Shareable PDF link (unguessable UUID) — sent to the party over WhatsApp."""
    inv = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
    if not inv:
        raise HTTPException(404, "Invoice not found")
    settings_doc = (await db.settings.find_one({"user_id": inv.get("user_id")}, {"_id": 0})) or Settings().model_dump()
    pdf = await _invoice_pdf_bytes(inv, settings_doc)
    filename = f"{_safe_filename(inv.get('party_name') or 'party')}_{inv.get('date','')}.pdf"
    return Response(pdf, media_type="application/pdf",
                    headers={"Content-Disposition": f"inline; filename={filename}"})


@api_router.get("/parties/{party_id}/export")
async def export_party_excel(party_id: str, user: dict = Depends(auth_mod.get_current_user_any)):
    uid = _uid(user)
    party = await db.parties.find_one({"id": party_id, "user_id": uid}, {"_id": 0})
    if not party:
        raise HTTPException(404, "Party not found")
    invoices = await db.invoices.find({"party_id": party_id, "user_id": uid}, {"_id": 0}).sort("date", -1).to_list(1000)
    settings_doc = await _get_settings_doc(uid)

    wb = Workbook()
    first = True
    if not invoices:
        ws = wb.active
        ws.title = "Empty"
        ws["A1"] = "No invoices yet"
    for inv in invoices:
        if first:
            ws = wb.active
            ws.title = (inv.get("date", "Invoice") or "Invoice")[:31]
            first = False
        else:
            ws = wb.create_sheet(title=(inv.get("date", "Invoice") or "Invoice")[:31])
        _build_invoice_sheet(ws, inv, settings_doc)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    safe_party = _safe_filename(party.get("name") or "party")
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={safe_party}_invoices.xlsx"},
    )


@public_router.get("/")
async def root():
    return {"message": "Commission Invoice API"}


# ---------------- GOOGLE SYNC WIRING ----------------
def _generate_invoice_bytes(inv: dict, settings_doc: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = inv.get("date", "Invoice")[:31]
    _build_invoice_sheet(ws, inv, settings_doc)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _sync_invoice_to_google(uid: str, invoice_id: str):
    try:
        if not await google_sync.is_connected(db, uid):
            return
        inv = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
        if not inv:
            return
        settings_doc = await _get_settings_doc(uid)
        xlsx = _generate_invoice_bytes(inv, settings_doc)
        safe_party = _safe_filename(inv.get("party_name") or "party")
        base = f"{safe_party}_{inv.get('date','')}"
        await google_sync.upload_invoice_file(db, uid, inv, xlsx, base + ".xlsx", "xlsx")
        try:
            pdf = await _invoice_pdf_bytes(inv, settings_doc)
            await google_sync.upload_invoice_file(db, uid, inv, pdf, base + ".pdf", "pdf")
        except Exception as e:
            logger.warning("pdf drive upload failed: %s", e)
        # re-read so sheet row carries the drive urls
        fresh = await db.invoices.find_one({"id": invoice_id}, {"_id": 0}) or inv
        await google_sync.upsert_master_sheet_row(db, uid, fresh)
        await google_sync.sync_ledger(db, uid)
    except Exception as e:
        logger.warning("google sync failed for %s: %s", invoice_id, e)


@api_router.post("/invoices/{invoice_id}/sync")
async def manual_sync_invoice(invoice_id: str, user: dict = Depends(get_current_user)):
    """Force a re-sync of this invoice to Drive + master Sheet."""
    uid = _uid(user)
    await _find_invoice(invoice_id, user)
    if not await google_sync.is_connected(db, uid):
        raise HTTPException(400, "Google account not connected")
    await _sync_invoice_to_google(uid, invoice_id)
    fresh = await db.invoices.find_one({"id": invoice_id}, {"_id": 0})
    return {"ok": True, "drive_file_url": fresh.get("drive_file_url"),
            "drive_pdf_url": fresh.get("drive_pdf_url"), "synced_at": fresh.get("synced_at")}


@api_router.post("/google/sync-all")
async def sync_all_invoices(user: dict = Depends(get_current_user)):
    """Push every existing invoice to Drive + Sheets (idempotent)."""
    uid = _uid(user)
    if not await google_sync.is_connected(db, uid):
        raise HTTPException(400, "Google account not connected")
    ids = await db.invoices.find({"user_id": uid}, {"_id": 0, "id": 1}).to_list(10000)
    done = 0
    for d in ids:
        try:
            await _sync_invoice_to_google(uid, d["id"])
            done += 1
        except Exception as e:
            logger.warning("sync-all: %s failed: %s", d.get("id"), e)
    return {"ok": True, "synced": done, "total": len(ids)}


app.include_router(api_router)
app.include_router(public_router)
app.include_router(auth_mod.router)
app.include_router(storage_mod.router)
google_sync.register_oauth(db, app)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(name)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


@app.on_event("startup")
async def run_migrations():
    try:
        await migrate_payments()
        await auth_mod.ensure_indexes()
    except Exception as e:
        logger.warning("startup migration failed: %s", e)
    try:
        await db.invoices.create_index(
            [("user_id", 1), ("import_source_key", 1)],
            unique=True,
            partialFilterExpression={"import_source_key": {"$type": "string"}},
        )
    except Exception as e:
        logger.warning("invoice import index setup failed: %s", e)
    try:
        await storage_mod.run_in_threadpool(storage_mod.init_storage)
    except Exception as e:
        logger.warning("object storage init failed: %s", e)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
