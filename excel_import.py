"""Parse party-wise commission invoice workbooks into normalized invoice data."""

from datetime import date, datetime
from io import BytesIO
import re
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

MAX_WORKBOOK_SHEETS = 500
MAX_WORKBOOK_ROWS = 250_000
MAX_WORKBOOK_CELLS = 2_000_000
MAX_INVOICES = 20_000
MAX_WARNINGS = 100


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _number(value: Any) -> float:
    if isinstance(value, (int, float)):
        return float(value)
    text = _text(value).replace(",", "").replace("₹", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _date(value: Any, epoch: datetime) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        converted = from_excel(value, epoch)
        if isinstance(converted, datetime):
            return converted.date().isoformat()
        if isinstance(converted, date):
            return converted.isoformat()
        raise ValueError(f"Unrecognized invoice date: {value}")

    text = _text(value)
    for fmt in ("%d/%m/%y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    raise ValueError(f"Unrecognized invoice date: {text or '(blank)'}")


def _find_header(
    rows: list[tuple[Any, ...]], start: int, end: int
) -> tuple[int, dict[str, int]] | None:
    required = {"SR", "ITEM", "QTY", "TOTAL"}
    for row_index in range(start, end):
        columns = {
            _text(value).upper(): idx
            for idx, value in enumerate(rows[row_index])
            if _text(value)
        }
        if required.issubset(columns):
            if "RATE" not in columns:
                columns["RATE"] = columns["QTY"] + 1
            if columns["RATE"] >= columns["TOTAL"]:
                continue
            return row_index, columns
    return None


def parse_commission_workbook(contents: bytes) -> tuple[list[dict], list[str]]:
    """Extract invoices from worksheets using the Commission Item workbook layout."""
    with ZipFile(BytesIO(contents)) as archive:
        members = archive.infolist()
        if "[Content_Types].xml" not in archive.namelist():
            raise ValueError("The selected file is not an .xlsx workbook")
        if (
            len(members) > 2_000
            or sum(member.file_size for member in members) > 100 * 1024 * 1024
        ):
            raise ValueError("Workbook expands beyond the safe import limit")

    workbook = load_workbook(BytesIO(contents), read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) > MAX_WORKBOOK_SHEETS:
            raise ValueError(
                f"Workbook has too many sheets (maximum {MAX_WORKBOOK_SHEETS})"
            )

        parsed: list[dict] = []
        warnings: list[str] = []
        row_count = 0
        cell_count = 0

        for sheet in workbook.worksheets:
            if sheet.max_row is not None:
                row_count += sheet.max_row
                cell_count += sheet.max_row * (sheet.max_column or 1)
            if row_count > MAX_WORKBOOK_ROWS or cell_count > MAX_WORKBOOK_CELLS:
                raise ValueError("Workbook is too large to import safely")

            rows = list(sheet.iter_rows(values_only=True))
            date_rows = [
                idx
                for idx, row in enumerate(rows)
                if any(_text(cell).casefold() == "date:" for cell in row)
            ]

            for position, date_index in enumerate(date_rows):
                next_index = (
                    date_rows[position + 1]
                    if position + 1 < len(date_rows)
                    else len(rows)
                )
                date_row = rows[date_index]
                party_name = _text(date_row[2] if len(date_row) > 2 else None)
                date_value = date_row[5] if len(date_row) > 5 else None
                source = f"{sheet.title}!{date_index + 1}"

                if not party_name and not _text(date_value):
                    continue
                try:
                    invoice_date = _date(date_value, workbook.epoch)
                except ValueError as exc:
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(f"{source}: {exc}; invoice skipped")
                    continue
                if not party_name:
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(
                            f"{source}: party name is blank; invoice skipped"
                        )
                    continue

                header = _find_header(rows, date_index + 1, next_index)
                if not header:
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(
                            f"{source}: item table header not found; invoice skipped"
                        )
                    continue

                header_index, columns = header
                total_index = next(
                    (
                        idx
                        for idx in range(header_index + 1, next_index)
                        if _text(
                            rows[idx][columns["RATE"]]
                            if len(rows[idx]) > columns["RATE"]
                            else None
                        ).casefold()
                        == "total"
                    ),
                    None,
                )
                if total_index is None:
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(
                            f"{source}: item total row not found; invoice skipped"
                        )
                    continue

                items = []
                for row in rows[header_index + 1 : total_index]:
                    item_name = _text(
                        row[columns["ITEM"]] if len(row) > columns["ITEM"] else None
                    )
                    qty = _number(
                        row[columns["QTY"]] if len(row) > columns["QTY"] else None
                    )
                    rate = _number(
                        row[columns["RATE"]] if len(row) > columns["RATE"] else None
                    )
                    total = _number(
                        row[columns["TOTAL"]] if len(row) > columns["TOTAL"] else None
                    )
                    if not total and qty and rate:
                        total = round(qty * rate, 2)
                    if not (item_name or qty or rate or total):
                        continue
                    sr_value = _number(
                        row[columns["SR"]] if len(row) > columns["SR"] else None
                    )
                    items.append(
                        {
                            "sr": int(sr_value) if sr_value.is_integer() else None,
                            "item": item_name,
                            "qty": qty,
                            "rate": rate,
                            "total": total,
                        }
                    )

                line_total = round(sum(item["total"] for item in items), 2)
                sheet_total = _number(
                    rows[total_index][columns["TOTAL"]]
                    if len(rows[total_index]) > columns["TOTAL"]
                    else None
                )
                if (
                    abs(sheet_total - line_total) > 0.01
                    and len(warnings) < MAX_WARNINGS
                ):
                    warnings.append(
                        f"{source}: line items total differs from the sheet total"
                    )

                sr_value = _number(
                    rows[total_index][columns["SR"]]
                    if len(rows[total_index]) > columns["SR"]
                    else None
                )
                bhada = 0.0
                mazdoori = 0.0
                commission_percent = 6.0
                commission_value = None
                sheet_net = None

                for row in rows[total_index + 1 : next_index]:
                    for idx, value in enumerate(row):
                        label = _text(value)
                        normalized = label.casefold()
                        following = next(
                            (_text(v) for v in row[idx + 1 :] if _text(v)), None
                        )
                        if "भाड़ा" in label or "bhada" in normalized:
                            bhada = _number(following)
                        elif "मजदूरी" in label or "mazdoori" in normalized:
                            mazdoori = _number(following)
                        elif label.strip().startswith("कमीशन") or normalized.startswith(
                            "commission"
                        ):
                            commission_value = _number(following)
                            match = re.search(r"(\d+(?:\.\d+)?)\s*%", label)
                            if match:
                                commission_percent = float(match.group(1))
                        elif "खर्च" in label or "kharch" in normalized:
                            pass
                        elif label == "नेट" or normalized == "net":
                            sheet_net = _number(following)

                calculated_commission = round(line_total * commission_percent / 100, 2)
                calculated_net = round(
                    line_total - bhada - mazdoori - calculated_commission, 2
                )
                if (
                    commission_value is not None
                    and abs(commission_value - calculated_commission) > 0.01
                ):
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(
                            f"{source}: commission mismatch "
                            f"({commission_value:g} stored; "
                            f"{calculated_commission:g} calculated)"
                        )
                if sheet_net is not None and abs(sheet_net - calculated_net) > 0.01:
                    if len(warnings) < MAX_WARNINGS:
                        warnings.append(
                            f"{source}: net amount differs from the app calculation"
                        )

                parsed.append(
                    {
                        "sheet_name": sheet.title,
                        "source_row": date_index + 1,
                        "party_name": party_name,
                        "date": invoice_date,
                        "items": items,
                        "sr_total": sr_value,
                        "bhada": bhada,
                        "mazdoori": mazdoori,
                        "commission_percent": commission_percent,
                    }
                )
                if len(parsed) > MAX_INVOICES:
                    raise ValueError(
                        f"Workbook has too many invoices (maximum {MAX_INVOICES})"
                    )

        return parsed, warnings
    finally:
        workbook.close()
